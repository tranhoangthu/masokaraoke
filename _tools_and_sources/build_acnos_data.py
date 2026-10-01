import os
import sys
import re
import json
import time
import openpyxl

if sys.stdout and sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def infer_acnos_genre(title):
    t_lower = title.lower()
    if 'remix' in t_lower or '(remix)' in t_lower or 'dj' in t_lower or 'dance' in t_lower:
        return 1  # Nhạc Trẻ / Remix
    if any(k in t_lower for k in ['(vc)', 'vọng cổ', '(tân cổ)', 'tân cổ', 'cổ nhạc', 'trích đoạn', 'lý ']):
        return 3  # Quê Hương - Cổ Nhạc
    if any(k in t_lower for k in ['thiếu nhi', 'chú ếch con', 'con cò bé bé', 'mẹ ơi']):
        return 5  # Thiếu Nhi
    return 0  # Khác / Pop

def clean_author_field(nhac, loi):
    nhac = (nhac or '').strip()
    loi = (loi or '').strip()
    if nhac in ['-', 'Music: -', 'Nhạc: -']: nhac = ''
    if loi in ['-', 'Lyrics: -', 'Lời: -']: loi = ''
    
    if nhac and loi and nhac.lower() != loi.lower():
        # Avoid duplicate combinations like "Trịnh Công Sơn, Trịnh Công Sơn"
        return f"{nhac}, {loi}"
    elif nhac:
        return nhac
    elif loi:
        return loi
    return ''

def build_acnos_dataset(excel_path='List ACNOS (2017).xlsx'):
    print("=" * 65)
    print("  XÂY DỰNG BỘ DỮ LIỆU ACNOS VOL 62 BẢN ĐẦY ĐỦ (2017)")
    print("=" * 65)
    
    t0 = time.time()
    
    # 1. Load previous acnos vols if available to preserve vol 58, 60
    old_vol_map = {}
    if os.path.exists('data/acnos.json'):
        try:
            with open('data/acnos.json', 'r', encoding='utf-8') as f:
                old_data = json.load(f)
                for r in old_data:
                    m = r[1]
                    v = r[7]
                    old_vol_map[m] = v
            print(f"Loaded {len(old_vol_map):,} historical song vols from old data/acnos.json")
        except Exception as e:
            print(f"Could not load old acnos.json: {e}")

    # 2. Read Excel
    print(f"Đọc file Excel: {excel_path} ...")
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    
    all_songs = []
    seen_keys = set()
    
    # Sheet 1: Tiếng Việt
    ws_vi = wb['Tiếng Việt']
    vi_count = 0
    for r in ws_vi.iter_rows(values_only=True):
        if not r or not r[0]: continue
        c0 = str(r[0]).strip()
        if c0 in ['MÃ SỐ', 'CODE'] or 'SONGBOOK' in c0: continue
        
        t_raw = str(r[1] or '').strip()
        a_raw = str(r[2] or '').strip()
        
        lines = [ln.strip() for ln in t_raw.split('\n') if ln.strip()]
        title = lines[0] if lines else ''
        intro = lines[1] if len(lines) > 1 else ''
        
        nhac = ''
        loi = ''
        for al in a_raw.split('\n'):
            al = al.strip()
            if al.lower().startswith('nhạc:') or al.lower().startswith('music:'):
                nhac = al.split(':', 1)[1].strip()
            elif al.lower().startswith('lời:') or al.lower().startswith('lyrics:'):
                loi = al.split(':', 1)[1].strip()
                
        author = clean_author_field(nhac, loi)
        
        # Parse code
        if '/' in c0:
            parts = [p.strip() for p in c0.split('/')]
            maso6_str = parts[0]
            try:
                maso = int(parts[1])
            except:
                maso = int(parts[0]) if parts[0].isdigit() else 0
        else:
            try:
                maso = int(c0)
            except:
                maso = 0
            maso6_str = c0 if len(c0) == 6 else ''
            
        genre = infer_acnos_genre(title)
        
        # Vol determination
        vol = old_vol_map.get(maso, 62)
        
        # Format intro with HDMI code if different
        if maso6_str and maso6_str != str(maso):
            intro_display = f"{intro} [Mã HDMI: {maso6_str}]".strip()
        else:
            intro_display = intro
            
        key = (maso, maso6_str, title.lower())
        if key not in seen_keys:
            seen_keys.add(key)
            all_songs.append({
                'maso': maso,
                'title': title,
                'intro': intro_display,
                'author': author,
                'nhac': nhac,
                'loi': loi,
                'lan': 0, # Tiếng Việt
                'g': genre,
                'v': vol,
                'maso6': maso6_str
            })
            vi_count += 1
            
    print(f"  -> Trích xuất thành công {vi_count:,} bài Tiếng Việt")
    
    # Sheet 2: Tiếng Anh
    ws_en = wb['Tiếng Anh']
    en_count = 0
    for r in ws_en.iter_rows(values_only=True):
        if not r or not r[0]: continue
        c0 = str(r[0]).strip()
        if c0 in ['MÃ SỐ', 'CODE'] or 'SONGBOOK' in c0: continue
        
        t_raw = str(r[1] or '').strip()
        a_raw = str(r[2] or '').strip()
        
        lines = [ln.strip() for ln in t_raw.split('\n') if ln.strip()]
        title = lines[0] if lines else ''
        intro = lines[1] if len(lines) > 1 else ''
        
        nhac = ''
        loi = ''
        for al in a_raw.split('\n'):
            al = al.strip()
            if al.lower().startswith('nhạc:') or al.lower().startswith('music:'):
                nhac = al.split(':', 1)[1].strip()
            elif al.lower().startswith('lời:') or al.lower().startswith('lyrics:'):
                loi = al.split(':', 1)[1].strip()
                
        author = clean_author_field(nhac, loi)
        
        try:
            maso = int(c0)
        except:
            maso = 0
        maso6_str = c0
        
        genre = infer_acnos_genre(title)
        vol = old_vol_map.get(maso, 62)
        
        if maso6_str and maso6_str != str(maso):
            intro_display = f"{intro} [Mã HDMI: {maso6_str}]".strip()
        else:
            intro_display = intro
            
        key = (maso, maso6_str, title.lower())
        if key not in seen_keys:
            seen_keys.add(key)
            all_songs.append({
                'maso': maso,
                'title': title,
                'intro': intro_display,
                'author': author,
                'nhac': nhac,
                'loi': loi,
                'lan': 1, # Tiếng Anh
                'g': genre,
                'v': vol,
                'maso6': maso6_str
            })
            en_count += 1
            
    print(f"  -> Trích xuất thành công {en_count:,} bài Tiếng Anh")
    print(f"Tổng cộng bài hát Acnos: {len(all_songs):,} bài (thời gian đọc: {time.time()-t0:.2f}s)")
    
    # Sort by maso
    all_songs.sort(key=lambda x: (x['maso'], x['title']))
    
    # Format rows: [id, maso, title, intro, author, lan, g, v]
    acnos_rows = []
    for i, s in enumerate(all_songs, start=1):
        row = [i, s['maso'], s['title'], s['intro'], s['author'], s['lan'], s['g'], s['v']]
        acnos_rows.append(row)
        
    # Write data/acnos.json and data/acnos.js
    os.makedirs('data', exist_ok=True)
    with open('data/acnos.json', 'w', encoding='utf-8') as f:
        json.dump(acnos_rows, f, ensure_ascii=False, separators=(',', ':'))
        
    with open('data/acnos.js', 'w', encoding='utf-8') as f:
        f.write("window.KARAOKE_DATA_ACNOS = ")
        json.dump(acnos_rows, f, ensure_ascii=False, separators=(',', ':'))
        f.write(";\n")
        
    print(f"Đã lưu data/acnos.json ({os.path.getsize('data/acnos.json'):,} bytes)")
    print(f"Đã lưu data/acnos.js ({os.path.getsize('data/acnos.js'):,} bytes)")
    
    # Update version.json
    if os.path.exists('version.json'):
        with open('version.json', 'r', encoding='utf-8') as f:
            vdata = json.load(f)
        vdata['datasets']['acnos'] = {
            'vol': 'Vol 62',
            'count': len(acnos_rows)
        }
        # Recalculate total songs in version.json
        total_songs = sum(d.get('count', 0) for d in vdata.get('datasets', {}).values())
        vdata['totalSongs'] = total_songs
        with open('version.json', 'w', encoding='utf-8') as f:
            json.dump(vdata, f, ensure_ascii=False, indent=2)
        print(f"Đã cập nhật version.json: acnos={len(acnos_rows):,} bài, tổng số bài toàn hệ thống={total_songs:,}")
        
    return all_songs

if __name__ == '__main__':
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    xlsx_path = os.path.join(base_dir, 'List ACNOS (2017).xlsx')
    build_acnos_dataset(xlsx_path)
