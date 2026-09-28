import sqlite3
import zipfile
import glob
import json
import os
import sys
import openpyxl
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

os.makedirs('data', exist_ok=True)
os.makedirs('data/lyrics', exist_ok=True)
os.makedirs('icons', exist_ok=True)

# 1. Read existing APK (Mã số Karaoke Vietnam 2.1.6)
apk_path = glob.glob('*Vietnam_2.1.6*.apk')[0]
print(f"Reading original APK: {apk_path}")
with zipfile.ZipFile(apk_path, 'r') as z:
    with open('temp_db.sqlite', 'wb') as f:
        f.write(z.read('assets/kokv1'))

conn = sqlite3.connect('temp_db.sqlite')
cur = conn.cursor()

# 1A. Arirang (a)
cur.execute("SELECT id, maso, title, intro, author, lan, g, v FROM a ORDER BY maso ASC")
arirang_rows = cur.fetchall()
print(f"Arirang: {len(arirang_rows)} songs")

# 1B. California (c)
cur.execute("SELECT id, maso, title, intro, author, lan, g, v FROM c ORDER BY maso ASC")
california_rows = cur.fetchall()
print(f"California: {len(california_rows)} songs")

# 1C. Việt KTV (v)
cur.execute("SELECT id, maso, title, intro, author, lan, g, v FROM v ORDER BY maso ASC")
vietktv_rows = cur.fetchall()
print(f"Việt KTV: {len(vietktv_rows)} songs")

# 1D. MusicCore (m) - merged with MusicCore Vol 102 Excel
cur.execute("SELECT id, maso, title, intro, author, lan, g, v FROM m ORDER BY maso ASC")
raw_mc = cur.fetchall()
mc_dict = {r[1]: list(r) for r in raw_mc}
max_mc_id = max(r[0] for r in raw_mc) if raw_mc else 0

# Load MusicCore Vol 102 Excel
mc_excel = 'Danh mục bài hát MusicCore (cập nhật tới vol 102).xlsx'
if os.path.exists(mc_excel):
    print(f"Merging {mc_excel}...")
    wb = openpyxl.load_workbook(mc_excel, data_only=True)
    sheet = wb.active
    added_mc = 0
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not row or row[0] is None: continue
        maso = int(float(row[0]))
        raw_title = row[1]
        if hasattr(raw_title, 'strftime'):
            title = raw_title.strftime('%H:%M').lstrip('0')
        else:
            title = str(raw_title or '').strip()
        composer = str(row[2] or '').strip()
        writer = str(row[3] or '').strip()
        author = composer if composer else writer
        
        if maso not in mc_dict:
            max_mc_id += 1
            # [id, maso, title, intro, author, lan, g, v]
            mc_dict[maso] = [max_mc_id, maso, title, '', author, 0, 0, 102]
            added_mc += 1
    wb.close()
    print(f"Added {added_mc} new songs from MusicCore Vol 102!")

musiccore_rows = sorted(mc_dict.values(), key=lambda x: x[1])
print(f"Total MusicCore songs: {len(musiccore_rows)}")

conn.close()
os.remove('temp_db.sqlite')

# 2. Paramax (Vietnamese & English from Excel files)
paramax_rows = []
paramax_id = 0

# 2A. Paramax VN
p_vn_file = 'Danh mục bài hát Paramax (cập nhật tới vol 52).xlsx'
if os.path.exists(p_vn_file):
    print(f"Reading {p_vn_file}...")
    wb = openpyxl.load_workbook(p_vn_file, data_only=True)
    sheet = wb.active
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not row or row[0] is None: continue
        maso = int(float(row[0]))
        title = str(row[1] or '').strip()
        singer = str(row[2] or '').strip()
        composer = str(row[3] or '').strip()
        lyric_author = str(row[4] or '').strip()
        intro = str(row[5] or '').strip()
        author = composer if composer else (singer if singer else lyric_author)
        if singer and singer != '-' and singer != author:
            author = f"{author} ({singer})" if author else singer
        paramax_id += 1
        paramax_rows.append([paramax_id, maso, title, intro, author, 0, 0, 52])
    wb.close()
    print(f"Paramax VN: {paramax_id} songs")

# 2B. Paramax EN
p_en_file = 'Danh mục bài hát tiếng Anh (Paramax) (cập nhật tới vol 52).xlsx'
if os.path.exists(p_en_file):
    print(f"Reading {p_en_file}...")
    wb = openpyxl.load_workbook(p_en_file, data_only=True)
    sheet = wb.active
    count_en = 0
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0 or not row or row[0] is None: continue
        maso = int(float(row[0]))
        title = str(row[1] or '').strip()
        singer = str(row[2] or '').strip()
        composer = str(row[3] or '').strip()
        intro = str(row[5] or '').strip()
        author = composer if composer else singer
        if singer and singer != '-' and singer != author:
            author = f"{author} ({singer})" if author else singer
        paramax_id += 1
        count_en += 1
        paramax_rows.append([paramax_id, maso, title, intro, author, 1, 0, 52])
    wb.close()
    print(f"Paramax EN: {count_en} songs")

paramax_rows.sort(key=lambda x: x[1])
print(f"Total Paramax songs: {len(paramax_rows)}")

# 3. Vitek VTB (from Vitek List APK)
vitek_apk = 'Vitek List_3.25062020.apk'
vitek_rows = []
if os.path.exists(vitek_apk):
    print(f"Reading Vitek APK: {vitek_apk}...")
    with zipfile.ZipFile(vitek_apk, 'r') as z:
        vitek_dict = {}
        for db_name in ['assets/vitek.sqlite', 'assets/vitek1.sqlite']:
            raw = z.read(db_name)
            decoded = bytes(b ^ 0x81 for b in raw)
            open('t_vitek.db', 'wb').write(decoded)
            c = sqlite3.connect('t_vitek.db')
            cur_v = c.cursor()
            for r in cur_v.execute('select number, title, short_lyric, music_by, singer, language, genre, version from song'):
                if not r[0]: continue
                try:
                    num = int(r[0])
                except:
                    continue
                if num not in vitek_dict:
                    title = str(r[1] or '').strip()
                    intro = str(r[2] or '').strip()
                    composer = str(r[3] or '').strip()
                    singer = str(r[4] or '').strip()
                    author = composer if composer and composer != '-' else singer
                    if singer and singer != '-' and singer != author:
                        author = f"{author} ({singer})" if author else singer
                    lang = 1 if r[5] == 'en' else 0
                    genre_str = str(r[6] or '').lower()
                    g = 1 if 'trẻ' in genre_str else (2 if 'tình' in genre_str or 'quê' in genre_str else 0)
                    ver = int(r[7]) if r[7] else 0
                    vitek_dict[num] = (title, intro, author, lang, g, ver)
            c.close()
            os.remove('t_vitek.db')
            
        v_id = 0
        for num in sorted(vitek_dict.keys()):
            v_id += 1
            t, intro, a, lang, g, ver = vitek_dict[num]
            vitek_rows.append([v_id, num, t, intro, a, lang, g, ver])
            
    print(f"Total Vitek VTB songs: {len(vitek_rows)}")

# 4. Save all datasets
all_datasets = {
    'a': ('arirang', arirang_rows, 'KARAOKE_DATA_A'),
    'm': ('musiccore', musiccore_rows, 'KARAOKE_DATA_M'),
    'c': ('california', california_rows, 'KARAOKE_DATA_C'),
    'v': ('vietktv', vietktv_rows, 'KARAOKE_DATA_V'),
    'p': ('paramax', paramax_rows, 'KARAOKE_DATA_P'),
    't': ('vitek', vitek_rows, 'KARAOKE_DATA_T')
}

total_all = 0
for key, (fname, rows, var_name) in all_datasets.items():
    total_all += len(rows)
    print(f"Writing data/{fname}.js and data/{fname}.json ({len(rows):,} songs)...")
    with open(f'data/{fname}.json', 'w', encoding='utf-8') as f:
        json.dump(rows, f, ensure_ascii=False, separators=(',', ':'))
    with open(f'data/{fname}.js', 'w', encoding='utf-8') as f:
        f.write(f"window.{var_name} = ")
        json.dump(rows, f, ensure_ascii=False, separators=(',', ':'))
        f.write(";")

print(f"\n==========================================")
print(f"SUCCESS! Total songs in all 6 systems: {total_all:,}")
print(f"==========================================")
